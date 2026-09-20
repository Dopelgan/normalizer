"""
Уровень 2 — извлечение текста оттуда, где он уже есть.

Три источника, объединённые одним свойством: содержимое лежит в файле как
текст, и распознавание не нужно. Отсюда и стоимость уровня — низкая, и
точность — полная.

Важная оговорка про PDF. Текстовый слой описывает текст, но ничего не знает
про таблицы, формулы и чертежи. Поэтому для документа со структурой этот
уровень намеренно получает низкую оценку структурной полноты и лестница
поднимается выше — к разбору с детекцией областей. Для простого текстового
PDF уровень 2 остаётся последним, и это экономит весь тяжёлый конвейер.
"""

from __future__ import annotations

import logging
import re
from typing import List, Optional

from core import filetypes
from core.ladder.base import Strategy
from core.ladder.context import DocumentContext
from core.models.parse_result import ParsedBlock, ParseResult
from core.providers.text_layer import group_lines, lines_to_text, merge_line_bboxes

logger = logging.getLogger(__name__)

# Строка-разделитель шапки таблицы GFM: |---|:---:|---:|
_DELIMITER_ROW = re.compile(r"^\s*\|?(?:\s*:?-{1,}:?\s*\|)+\s*:?-{1,}:?\s*\|?\s*$")

try:
    import docx as python_docx
    DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover
    python_docx = None
    DOCX_AVAILABLE = False


class PdfTextLayerStrategy(Strategy):
    """Текстовый слой векторного PDF без обращения к тяжёлому парсеру."""

    level = 2
    name = "pdf_text_layer"
    method = "text_layer"
    exhaustive = True

    def applicable(self, context: DocumentContext) -> bool:
        return context.kind == filetypes.KIND_PDF and context.text_layer is not None

    def run(self, context: DocumentContext) -> ParseResult:
        layer = context.text_layer
        if layer is None:
            raise ValueError("У документа нет текстового слоя")

        blocks: List[ParsedBlock] = []
        for page in sorted(layer.pages):
            lines = [line for line in layer.lines(page) if line.char_count()]
            for group in group_lines(lines):
                text = lines_to_text(group)
                if not text:
                    continue
                blocks.append(ParsedBlock(
                    type="text", text=text, page=page,
                    bbox=merge_line_bboxes(group),
                    confidence=1.0, method=self.method,
                ))

        if not blocks:
            raise ValueError("Текстовый слой пуст")

        return self.build_result(
            blocks, source_kind="vector_pdf", page_count=layer.page_count,
            text_layer_chars=layer.char_counts_by_page(),
        )


class DocxStrategy(Strategy):
    """
    DOCX: абзацы и таблицы берутся из разметки документа как есть.

    DOC (Word 97-2003) обрабатывается тем же уровнем: он сначала
    конвертируется в DOCX, потому что плоский текст из двоичного формата
    теряет таблицы, а таблицы в этих документах и есть содержание.
    """

    level = 2
    name = "docx"
    method = "office_extraction"
    exhaustive = True

    def applicable(self, context: DocumentContext) -> bool:
        if not (DOCX_AVAILABLE and context.kind == filetypes.KIND_OFFICE_TEXT):
            return False
        from core.providers import office_convert

        # Формат, который без конвертера не открыть, уровню неприменим:
        # пусть лестница идёт выше, к распознаванию, а не падает здесь.
        if office_convert.converts(context.file_type):
            return office_convert.available()
        return True

    def run(self, context: DocumentContext) -> ParseResult:
        import io

        from core.providers import office_convert

        data = office_convert.convert(context.data, context.file_type)
        document = python_docx.Document(io.BytesIO(data))
        blocks: List[ParsedBlock] = []

        section_title: Optional[str] = None
        for paragraph in document.paragraphs:
            text = " ".join(paragraph.text.split())
            if not text:
                continue
            is_heading = (paragraph.style.name or "").lower().startswith("heading")
            if is_heading:
                section_title = text[:200]
            blocks.append(ParsedBlock(
                type="text", text=text, page=1, confidence=1.0,
                section_title=section_title, method=self.method,
            ))

        for table in document.tables:
            parsed = self._table(table)
            if parsed:
                blocks.append(ParsedBlock(
                    type="table", table_data=parsed, page=1, confidence=1.0,
                    section_title=section_title, method=self.method,
                ))

        if not blocks:
            raise ValueError("DOCX не содержит текста")

        # У DOCX нет постраничной разбивки без рендера: разбиение на
        # страницы делает текстовый процессор при печати, в файле его нет.
        return self.build_result(blocks, source_kind="office", page_count=1)

    @staticmethod
    def _table(table) -> Optional[dict]:
        rows = [[" ".join(cell.text.split()) for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            return None
        return {"headers": rows[0], "rows": rows[1:], "total_row": None}


class PlainTextStrategy(Strategy):
    """
    TXT и MD: абзацы разделяются пустой строкой, заголовки MD — секции,
    таблицы MD — таблицы.

    Про таблицы. Разметка в markdown задаёт структуру не хуже, чем ячейки
    XLSX: строка-разделитель `|---|---|` отличает таблицу от текста с
    вертикальными чертами однозначно. Раньше таблица уезжала в индекс
    плоским текстом вместе с палками и дефисами — по такому фрагменту не
    ответить ни на один вопрос про значение в ячейке. Разбор здесь ничего
    не угадывает: нет строки-разделителя — нет и таблицы.
    """

    level = 2
    name = "plain_text"
    method = "plain_text"
    exhaustive = True

    def applicable(self, context: DocumentContext) -> bool:
        return context.kind == filetypes.KIND_PLAIN_TEXT

    def run(self, context: DocumentContext) -> ParseResult:
        text = self._decode(context.data)
        if not text.strip():
            raise ValueError("Файл пуст")

        is_markdown = context.file_type == "md"
        blocks: List[ParsedBlock] = []
        section_title: Optional[str] = None

        for kind, payload in self._segments(text, is_markdown):
            if kind == "table":
                blocks.append(ParsedBlock(
                    type="table", page=1, confidence=1.0,
                    section_title=section_title, method=self.method,
                    table_data=payload,
                ))
                continue

            body = payload.strip("\n")
            if not body.strip():
                continue
            if is_markdown:
                heading = self._heading(body)
                if heading:
                    section_title = heading
            blocks.append(ParsedBlock(
                type="text", text=body.strip(), page=1, confidence=1.0,
                section_title=section_title, method=self.method,
            ))

        if not blocks:
            raise ValueError("Файл пуст")

        return self.build_result(blocks, source_kind="plain_text", page_count=1)

    # ------------------------------------------------------------ разбор
    @classmethod
    def _segments(cls, text: str, is_markdown: bool):
        """Куски документа по порядку: ('text', абзац) и ('table', данные)."""
        if not is_markdown:
            for chunk in text.split("\n\n"):
                yield "text", chunk
            return

        lines = text.split("\n")
        buffer: List[str] = []
        index = 0
        while index < len(lines):
            table, consumed = cls._read_table(lines, index)
            if table is not None:
                yield from cls._flush(buffer)
                buffer = []
                yield "table", table
                index += consumed
                continue
            buffer.append(lines[index])
            index += 1
        yield from cls._flush(buffer)

    @staticmethod
    def _flush(buffer: List[str]):
        for chunk in "\n".join(buffer).split("\n\n"):
            if chunk.strip():
                yield "text", chunk

    @classmethod
    def _read_table(cls, lines: List[str], start: int):
        """
        Таблица GFM, начиная со строки `start`. Возвращает (данные, сколько
        строк занято) или (None, 0). Признак таблицы — строка-разделитель
        сразу под шапкой, с тем же числом колонок.
        """
        if start + 1 >= len(lines):
            return None, 0
        header = cls._row(lines[start])
        if header is None:
            return None, 0
        if not _DELIMITER_ROW.match(lines[start + 1]):
            return None, 0
        delimiter = cls._row(lines[start + 1])
        if delimiter is None or len(delimiter) != len(header):
            return None, 0

        rows: List[List[str]] = []
        index = start + 2
        while index < len(lines):
            row = cls._row(lines[index])
            if row is None:
                break
            # Строки короче или длиннее шапки выравниваются: в живых файлах
            # это опечатка разметки, а не другая таблица.
            if len(row) < len(header):
                row = row + [""] * (len(header) - len(row))
            elif len(row) > len(header):
                row = row[:len(header)]
            rows.append(row)
            index += 1

        return (
            {"headers": header, "rows": rows, "total_row": None},
            index - start,
        )

    @staticmethod
    def _row(line: str) -> Optional[List[str]]:
        """Ячейки строки таблицы. `None` — строка таблицей не является."""
        stripped = line.strip()
        if "|" not in stripped:
            return None
        # Экранированная черта — это содержимое ячейки, а не разделитель.
        guarded = stripped.replace("\\|", "\x00")
        if guarded.startswith("|"):
            guarded = guarded[1:]
        if guarded.endswith("|"):
            guarded = guarded[:-1]
        cells = [cell.strip().replace("\x00", "|") for cell in guarded.split("|")]
        if len(cells) < 2:
            return None
        return cells

    @staticmethod
    def _decode(data: bytes) -> str:
        """UTF-8, иначе кодировки, в которых обычно приходят русские .txt."""
        for encoding in ("utf-8-sig", "utf-8", "cp1251", "koi8-r"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="replace")

    @staticmethod
    def _heading(block: str) -> Optional[str]:
        first = block.lstrip().split("\n", 1)[0]
        if first.startswith("#"):
            return first.lstrip("#").strip()[:200] or None
        return None
