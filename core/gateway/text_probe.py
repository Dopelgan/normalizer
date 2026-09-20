"""
Текст документа для классификации приёма.

Слой G-2 раньше искал маркеры в первых 64 КБ файла как есть. Для PDF, DOCX
и XLSX это сжатый двоичный поток: декодированный в cp1251, он даёт случайную
кириллицу и случайные же совпадения. Реальный случай: приложение к договору
было признано перепиской, потому что внутри сжатого потока нашлась
последовательность `re:`.

Здесь из файла достаётся именно текст — тем инструментом, который понимает
формат. Если текста нет (скан, картинка, чертёж), это честно сообщается:
классифицировать такой файл по содержимому нельзя, решают путь и имя.
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass
from typing import Optional

from core import filetypes

logger = logging.getLogger(__name__)

# Сколько текста достаточно, чтобы понять род документа. Первые страницы
# несут титул, шапку и предмет — дальше начинается тело, которое род уже
# не меняет.
TEXT_LIMIT = 20_000
PDF_PAGES = 5
XLSX_ROWS = 60

# Источник текста — попадает в сигналы приёма.
SOURCE_PDF = "pdf"
SOURCE_DOCX = "docx"
SOURCE_XLSX = "xlsx"
SOURCE_PLAIN = "plain"
SOURCE_NONE = "none"

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass(frozen=True)
class TextProbe:
    """Что удалось прочитать из файла и можно ли этому верить."""

    text: str
    source: str
    reason: Optional[str] = None

    @property
    def reliable(self) -> bool:
        """Текст извлечён инструментом формата и не пуст."""
        return self.source != SOURCE_NONE and bool(self.text.strip())

    def as_dict(self) -> dict:
        return {
            "text_source": self.source,
            "text_chars": len(self.text),
            **({"text_reason": self.reason} if self.reason else {}),
        }


def extract(file_type: str, data: bytes) -> TextProbe:
    """Текст документа по его типу. Байты — весь файл, а не образец."""
    kind = filetypes.kind_of(file_type)

    if kind == filetypes.KIND_PDF:
        return _from_pdf(data)
    if kind == filetypes.KIND_OFFICE_TEXT:
        return _from_docx(data)
    if file_type == "xlsx":
        return _from_xlsx(data)
    if kind in (filetypes.KIND_PLAIN_TEXT, filetypes.KIND_SPREADSHEET):
        return _from_plain(data)
    if kind == filetypes.KIND_CAD:
        # DXF — текст, но это коды и координаты, а не проза: маркеры по
        # нему искать бессмысленно, род и так известен из формата.
        return TextProbe("", SOURCE_NONE, "формат CAD: текста для разбора нет")
    if kind == filetypes.KIND_IMAGE:
        return TextProbe("", SOURCE_NONE, "растр: текста без распознавания нет")
    return TextProbe("", SOURCE_NONE, f"тип {file_type!r} без извлечения текста")


# ---------------------------------------------------------------- форматы

def _from_pdf(data: bytes) -> TextProbe:
    from core.providers.text_layer import open_pdf

    document = open_pdf(data)
    if document is None:
        return TextProbe("", SOURCE_NONE, "PDF не открылся")
    try:
        chunks = []
        for index in range(min(document.page_count, PDF_PAGES)):
            try:
                chunks.append(document[index].get_text() or "")
            except Exception as exc:  # noqa: BLE001 — битая страница не роняет разбор
                logger.debug("Страница %d не прочиталась: %s", index + 1, exc)
            if sum(len(c) for c in chunks) >= TEXT_LIMIT:
                break
        text = _clean("\n".join(chunks))
    finally:
        document.close()

    if not text.strip():
        # Скан: текстового слоя нет. Это не ошибка, но и судить по
        # содержимому не о чем — маркеры применять нельзя.
        return TextProbe("", SOURCE_NONE, "PDF без текстового слоя (скан)")
    return TextProbe(text, SOURCE_PDF)


def _from_docx(data: bytes) -> TextProbe:
    try:
        import docx  # type: ignore
    except ImportError:
        return TextProbe("", SOURCE_NONE, "python-docx недоступен")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        return TextProbe("", SOURCE_NONE, f"DOCX не открылся: {exc}")

    parts = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells if cell.text))
            if sum(len(p) for p in parts) >= TEXT_LIMIT:
                break
    text = _clean("\n".join(parts))
    if not text.strip():
        return TextProbe("", SOURCE_NONE, "DOCX без текста")
    return TextProbe(text, SOURCE_DOCX)


def _from_xlsx(data: bytes) -> TextProbe:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        return TextProbe("", SOURCE_NONE, "openpyxl недоступен")
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        return TextProbe("", SOURCE_NONE, f"XLSX не открылся: {exc}")

    parts = []
    try:
        for sheet in book.worksheets:
            parts.append(str(sheet.title))
            for row_no, row in enumerate(sheet.iter_rows(values_only=True)):
                parts.append(" ".join(str(v) for v in row if v is not None))
                if row_no + 1 >= XLSX_ROWS:
                    break
            if sum(len(p) for p in parts) >= TEXT_LIMIT:
                break
    finally:
        book.close()

    text = _clean("\n".join(parts))
    if not text.strip():
        return TextProbe("", SOURCE_NONE, "XLSX без содержимого")
    return TextProbe(text, SOURCE_XLSX)


def _from_plain(data: bytes) -> TextProbe:
    for encoding in ("utf-8", "cp1251"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover — cp1251 декодирует почти всё
        text = data.decode("utf-8", errors="ignore")

    text = _clean(text)
    if not text.strip():
        return TextProbe("", SOURCE_NONE, "файл пуст")
    return TextProbe(text, SOURCE_PLAIN)


# ------------------------------------------------------------- внутреннее

def _clean(text: str) -> str:
    """Убрать управляющие символы и подрезать до предела."""
    return _CONTROL.sub(" ", text)[:TEXT_LIMIT]
