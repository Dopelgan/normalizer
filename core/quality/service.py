"""
DOCUMENT QUALITY GATE — проверка технической пригодности.

Отвечает на вопрос «пригоден ли документ технически», и только на него.
Вопрос «относится ли это к корпоративным знаниям» решён раньше, на Data
Gateway: сюда файл приходит уже признанным знанием.

Быстрый OCR здесь нужен лишь для предварительной оценки читаемости —
полный разбор выполняется позже и только после прохождения проверки. Если
у документа есть текстовый слой, распознавать вообще нечего: читаемость
доказана самим файлом.

Пять этапов: техническая валидация, дубликаты, дата и актуальность,
качество распознавания, классификация и маршрутизация. Предложение
размещения формирует система — утверждает человек.
"""

from __future__ import annotations

import hashlib
import logging
import os
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

from core import filetypes
from core.config import settings
from core.gateway.profile import GatewayProfile, load_profile
from core.gateway.service import ACCEPT, QUARANTINE, REJECT
from core.ladder.context import (
    GENRE_DRAWING,
    GENRE_MIXED,
    GENRE_SPREADSHEET,
    GENRE_TEXT,
    SOURCE_RASTER,
    DocumentContext,
)
from core.providers.storage import StorageProvider, StorageProviderFactory

logger = logging.getLogger(__name__)

# Типы маршрутизации из QG-5.
ROUTE_TEXT = "текстовый"
ROUTE_TABULAR = "табличный"
ROUTE_MIXED = "смешанный"
ROUTE_MULTIMODAL = "мультимодальный"
ROUTE_DRAWING = "чертёж"

# Расстояние Хэмминга между отпечатками, ниже которого документы считаются
# почти одинаковыми. 64-битный отпечаток, 6 бит — это около 90% сходства.
_SIMHASH_BITS = 64
_NEAR_DUPLICATE_DISTANCE = 6

# Год в имени файла — признак слабый, поэтому он должен быть отдельным
# числом, а не куском слова: `деталь_2000шт.pdf` получал дату 1 января 2000
# года и предупреждение «документу 26 лет — он может быть неактуален».
_NAME_WORD = "0-9A-Za-zА-Яа-яЁё"
_DATE_IN_NAME = re.compile(
    rf"(?<![{_NAME_WORD}])(19|20)\d{{2}}(?:[-_.](?:0[1-9]|1[0-2]))?(?![{_NAME_WORD}])"
)

# Уверенность «не измерено»: OCR не запустился, текста в формате не нашлось.
# Это не оценка качества, а признак того, что вопрос отложен до разбора.
_UNMEASURED_CONFIDENCE = 0.5

# Проба текста у форматов, которые читаются напрямую. Меньше этого числа
# символов — содержимое, скорее всего, лежит картинками.
_MIN_NATIVE_CHARS = 20
_NATIVE_PROBE_LIMIT = 4000
_NATIVE_PROBE_ROWS = 50

# Сколько символов текста быстрого OCR достаточно для отпечатка документа.
_OCR_SAMPLE_LIMIT = 4000


@dataclass
class QualityVerdict:
    """Решение Quality Gate вместе с предупреждениями и предложением места."""

    outcome: str
    reason: str
    stage: str
    warnings: List[str] = field(default_factory=list)
    confidence: float = 1.0
    file_hash: Optional[str] = None
    routing: Dict[str, Any] = field(default_factory=dict)
    signals: Dict[str, Any] = field(default_factory=dict)

    @property
    def accepted(self) -> bool:
        return self.outcome == ACCEPT

    def as_dict(self) -> Dict[str, Any]:
        return {
            "outcome": self.outcome,
            "reason": self.reason,
            "stage": self.stage,
            "warnings": self.warnings,
            "confidence": round(self.confidence, 3),
            "file_hash": self.file_hash,
            "routing": self.routing,
            "signals": self.signals,
        }


class QualityGate:
    """Пять этапов проверки перед тем, как пускать файл в нормализацию."""

    def __init__(
        self,
        profile: Optional[GatewayProfile] = None,
        storage: Optional[StorageProvider] = None,
        known_hashes: Optional[Dict[str, str]] = None,
        known_fingerprints: Optional[List[Tuple[str, int]]] = None,
    ):
        self.profile = profile or load_profile()
        self.storage = storage or StorageProviderFactory.default()
        # Отпечатки и хеши уже принятого приходят снаружи: ходить в БД из
        # проверки не нужно, а тестировать так заметно проще.
        self.known_hashes = dict(known_hashes or {})
        self.known_fingerprints = list(known_fingerprints or [])
        # Проба последнего быстрого OCR: её текст — единственный источник
        # отпечатка у растрового документа, где текстового слоя нет вовсе.
        self._last_probe: Optional[Tuple[str, "_OcrProbe"]] = None

    # ------------------------------------------------------------- публичное
    def evaluate(self, s3_fileid: str, uri: Optional[str] = None) -> QualityVerdict:
        path = uri or s3_fileid
        warnings: List[str] = []

        # Размер проверяется до чтения: файл на гигабайты нельзя сперва
        # затащить целиком в память, а потом решить, что он слишком большой.
        declared = self._declared_size(path)
        if declared is not None and declared > settings.QG_MAX_SIZE_BYTES:
            return QualityVerdict(
                QUARANTINE,
                f"Файл больше предела ({declared} Б против "
                f"{settings.QG_MAX_SIZE_BYTES} Б) — нужно решение человека.",
                "QG-1",
            )

        try:
            data = self.storage.read_bytes(path)
        except Exception as exc:  # noqa: BLE001
            return QualityVerdict(
                REJECT, f"Файл не читается из хранилища: {exc}", "QG-1"
            )

        # Тип берётся по содержимому: дальше от него зависит и набор
        # проверок, и то, какой уровень разбора возьмётся за документ.
        resolved = filetypes.resolve(path, data)
        if resolved.mismatch:
            warnings.append(resolved.explanation)

        verdict = self.technical_validation(path, data, resolved)
        if verdict is not None:
            verdict.warnings = warnings + verdict.warnings
            return verdict

        file_hash = hashlib.sha256(data).hexdigest()

        context = DocumentContext(
            uri=path, file_type=resolved.file_type,
            storage=self.storage, data=data,
        )

        fingerprint = _simhash(_text_sample(context))
        duplicate = self.duplicate_check(
            s3_fileid, file_hash, context, warnings, fingerprint=fingerprint
        )
        if duplicate is not None:
            return duplicate

        self.freshness_check(path, data, warnings)

        readability = self.ocr_check(context, warnings)

        if not fingerprint:
            # У растра текста до распознавания нет, и поиск почти-дублей на
            # сканах не работал вовсе. Текст быстрого OCR к этому моменту уже
            # посчитан — отпечаток снимается с него и сравнивается здесь.
            fingerprint = _simhash(self._ocr_sample(context))
            if fingerprint:
                self._near_duplicate(s3_fileid, fingerprint, warnings)

        # Маршрутизация и сигналы считаются и для карантина. Прежде вердикт
        # QG-4 уходил без них, и в очереди карантина не было видно ни рода
        # документа, ни качества растра — того самого, по чему человек и
        # принимает решение.
        routing, signals = self._describe(context, fingerprint)

        if readability.outcome != ACCEPT:
            readability.file_hash = file_hash
            readability.warnings = warnings + readability.warnings
            readability.routing = routing
            readability.signals = signals
            return readability

        return QualityVerdict(
            outcome=ACCEPT,
            reason="Документ технически пригоден к нормализации.",
            stage="QG-5",
            warnings=warnings,
            confidence=readability.confidence,
            file_hash=file_hash,
            routing=routing,
            signals=signals,
        )

    def _declared_size(self, path: str) -> Optional[int]:
        """Размер по метаданным хранилища. None — хранилище его не отдало."""
        getter = getattr(self.storage, "size", None)
        if getter is None:
            return None
        try:
            return getter(path)
        except Exception as exc:  # noqa: BLE001 — неизвестный размер не отказ
            logger.warning("Размер %s не определён: %s", path, exc)
            return None

    # -------------------------------------------- QG-1 техническая валидация
    def technical_validation(
        self,
        path: str,
        data: bytes,
        resolved: Optional[filetypes.TypeVerdict] = None,
    ) -> Optional[QualityVerdict]:
        if resolved is None:
            resolved = filetypes.resolve(path, data)
        extension = resolved.file_type

        reason = filetypes.rejection_reason(extension)
        if reason:
            if resolved.mismatch:
                reason = f"{reason} {resolved.explanation}"
            return QualityVerdict(REJECT, reason, "QG-1")

        if len(data) > settings.QG_MAX_SIZE_BYTES:
            return QualityVerdict(
                QUARANTINE,
                f"Файл больше предела ({len(data)} Б против "
                f"{settings.QG_MAX_SIZE_BYTES} Б) — нужно решение человека.",
                "QG-1",
            )

        if len(data) < self.profile.min_size_bytes:
            return QualityVerdict(
                REJECT,
                f"Файл слишком мал ({len(data)} Б) — содержимого в нём нет.",
                "QG-1",
            )

        opened, detail = _can_open(data, extension)
        if not opened:
            return QualityVerdict(
                REJECT, f"Файл повреждён и не открывается: {detail}", "QG-1"
            )
        return None

    # -------------------------------------------- QG-2 проверка на дубликаты
    def duplicate_check(
        self,
        s3_fileid: str,
        file_hash: str,
        context: DocumentContext,
        warnings: List[str],
        fingerprint: Optional[int] = None,
    ) -> Optional[QualityVerdict]:
        """
        Точный дубль отклоняется. Семантическое сходство выше порога — это
        не отказ, а флаг администратору: «обновление или дубль».
        """
        twin = self.known_hashes.get(file_hash)
        if twin and twin != s3_fileid:
            return QualityVerdict(
                REJECT,
                f"Точный дубликат уже принятого файла {twin!r} "
                f"(совпадает контрольная сумма).",
                "QG-2", file_hash=file_hash,
            )

        if fingerprint is None:
            fingerprint = _simhash(_text_sample(context))
        if fingerprint:
            self._near_duplicate(s3_fileid, fingerprint, warnings)
        return None

    def _near_duplicate(
        self, s3_fileid: str, fingerprint: int, warnings: List[str]
    ) -> Optional[str]:
        """Ближайший почти-дубль среди принятого. Это флаг, а не отказ."""
        for other_id, other_print in self.known_fingerprints:
            if other_id == s3_fileid:
                continue
            distance = _hamming(fingerprint, other_print)
            if distance <= _NEAR_DUPLICATE_DISTANCE:
                similarity = 1 - distance / _SIMHASH_BITS
                if similarity >= settings.QG_SIMILARITY_THRESHOLD:
                    warnings.append(
                        f"Похож на {other_id!r} (сходство {similarity:.0%}) — "
                        f"обновление или дубль, нужно решение администратора."
                    )
                    return other_id
        return None

    # ------------------------------------------- QG-3 дата и актуальность
    def freshness_check(self, path: str, data: bytes, warnings: List[str]) -> None:
        created = _document_date(path, data)
        if created is None:
            warnings.append(
                "У документа не удалось определить дату — актуальность "
                "не подтверждена, уверенность снижена."
            )
            return

        age_days = (datetime.now(timezone.utc) - created).days
        if age_days > settings.QG_MAX_AGE_DAYS:
            warnings.append(
                f"Документу {age_days // 365} лет — он может быть неактуален."
            )

    # ------------------------------- QG-4 проверка качества распознавания
    def ocr_check(self, context: DocumentContext, warnings: List[str]) -> QualityVerdict:
        """
        Быстрый OCR на первых страницах. У векторного документа читаемость
        доказана текстовым слоем — распознавать нечего.
        """
        if context.kind in (
            filetypes.KIND_PLAIN_TEXT, filetypes.KIND_SPREADSHEET,
            filetypes.KIND_OFFICE_TEXT,
        ):
            return self._direct_read_check(context, warnings)

        if context.kind == filetypes.KIND_CAD:
            # У исходника САПР содержимое — геометрия, а не текст: мерить
            # его пробой текста нечем и незачем.
            return QualityVerdict(ACCEPT, "Содержимое читается напрямую.", "QG-4")

        layer = context.text_layer
        if layer is not None:
            return self._layer_check(layer, warnings)

        probe = self._quick_ocr(context)
        # Текст пробы — единственный источник отпечатка у растра, где
        # текстового слоя нет; он понадобится уже после этой проверки.
        self._last_probe = (context.uri, probe) if probe is not None else None
        if probe is None:
            warnings.append(
                "Быстрый OCR недоступен — читаемость не проверена, "
                "оценка отложена до полного разбора."
            )
            return QualityVerdict(ACCEPT, "Читаемость не проверена.", "QG-4",
                                  confidence=_UNMEASURED_CONFIDENCE)

        if not probe.blocks:
            return self._no_text_verdict(context, warnings)

        confidence = probe.confidence
        if confidence < settings.QG_OCR_BLOCK_CONFIDENCE:
            return QualityVerdict(
                QUARANTINE,
                f"Распознавание почти ничего не даёт (уверенность "
                f"{confidence:.2f}). Документ нечитаем в оригинале — "
                f"нужно решение человека, а не повторная попытка.",
                "QG-4", confidence=confidence,
            )

        # Среднее по листу вытягивается крупной шапкой: лист, где читается
        # заголовок и не читается тело, проходил приём без единого
        # предупреждения. Решает доля нечитаемого, а не одно число.
        if probe.unreadable_share > settings.QG_OCR_UNREADABLE_SHARE:
            return QualityVerdict(
                QUARANTINE,
                f"Ниже порога распознано {probe.unreadable_share:.0%} текста "
                f"листа при средней уверенности {confidence:.2f}: читается "
                f"только часть документа — нужно решение человека.",
                "QG-4", confidence=confidence,
            )
        if probe.unreadable_share:
            warnings.append(
                f"{probe.unreadable_share:.0%} распознанного текста ниже порога "
                f"читаемости — эта часть содержимого может быть потеряна."
            )
        if confidence < settings.QG_OCR_WARN_CONFIDENCE:
            warnings.append(
                f"Низкое качество распознавания (уверенность {confidence:.2f}): "
                f"часть содержимого может быть потеряна."
            )
        return QualityVerdict(ACCEPT, "Читаемость подтверждена.", "QG-4",
                              confidence=confidence)

    @staticmethod
    def _layer_check(layer, warnings: List[str]) -> QualityVerdict:
        """
        Текстовый слой доказывает читаемость только тех страниц, которые он
        покрывает.

        Слой считается пригодным уже при половине страниц с текстом
        (`TextLayer.is_usable`) — у гибридного PDF, где часть листов
        вставлена сканами, это верно. Но вердикт при этом объявлял читаемым
        весь документ с уверенностью 1.0, и вторая половина уезжала в
        разбор непроверенной, без единого слова в отчёте приёма.
        """
        total = layer.page_count or 0
        covered = layer.pages_with_text()
        if total and covered < total:
            share = round(covered / total, 3)
            warnings.append(
                f"Текстовый слой покрывает {covered} страниц из {total}: "
                f"остальные — изображения, их читаемость проверит разбор."
            )
            return QualityVerdict(
                ACCEPT, "Текстовый слой покрывает документ частично.",
                "QG-4", confidence=share,
            )
        return QualityVerdict(
            ACCEPT, "У документа есть текстовый слой — читаемость доказана.",
            "QG-4", confidence=1.0,
        )

    @staticmethod
    def _no_text_verdict(
        context: DocumentContext, warnings: List[str]
    ) -> QualityVerdict:
        """
        Быстрый OCR не нашёл ни строки.

        Раньше пустой результат превращался в уверенность 0.0 и карантин с
        формулировкой «документ нечитаем». Но «текста нет» и «текст есть и
        не прочитан» — разные вещи: на чертеже или схеме надписей может не
        быть вовсе, и такой лист уходил человеку по ложной тревоге. Там,
        где текст ожидается — скан страницы, фотокопия, снимок экрана, —
        пустой результат остаётся поводом для карантина.
        """
        if context.classification.genre == GENRE_DRAWING:
            warnings.append(
                "Надписей на листе не найдено — содержимое оценит разбор чертежа."
            )
            return QualityVerdict(
                ACCEPT,
                "Текста на листе нет: это чертёж, а не страница текста.",
                "QG-4", confidence=_UNMEASURED_CONFIDENCE,
            )
        return QualityVerdict(
            QUARANTINE,
            "Распознавание не дало ни строки текста. Документ нечитаем в "
            "оригинале — нужно решение человека, а не повторная попытка.",
            "QG-4", confidence=0.0,
        )

    def _direct_read_check(
        self, context: DocumentContext, warnings: List[str]
    ) -> QualityVerdict:
        """
        Форматы с извлекаемым содержимым.

        «Читается напрямую» было утверждением о формате, а не о файле: DOCX
        из одних картинок и книга с пустым листом проходили приём с
        уверенностью 1.0 и доезжали до разбора, где выяснялось, что читать
        нечего. Проба текста стоит дешевле разбора.
        """
        text = _native_text(context)
        if text is None:
            return QualityVerdict(ACCEPT, "Содержимое читается напрямую.", "QG-4")
        if len(text.strip()) < _MIN_NATIVE_CHARS:
            warnings.append(
                "Извлекаемого текста в файле почти нет — содержимое, скорее "
                "всего, лежит картинками; читаемость оценит разбор."
            )
            return QualityVerdict(
                ACCEPT, "Извлекаемого текста в файле почти нет.", "QG-4",
                confidence=_UNMEASURED_CONFIDENCE,
            )
        return QualityVerdict(ACCEPT, "Содержимое читается напрямую.", "QG-4")

    def _quick_ocr(self, context: DocumentContext) -> Optional["_OcrProbe"]:
        """
        Быстрый OCR первых страниц. `None` — распознать не удалось вовсе;
        пустая проба — распознали, но текста не нашли. Разные случаи.
        """
        from core.providers.tesseract_fallback import (
            OcrUnavailable,
            TesseractFallbackProvider,
        )

        try:
            pages = list(range(1, max(1, settings.QG_OCR_PAGES) + 1))
            blocks = TesseractFallbackProvider(storage=self.storage).parse_pages(
                # Тип берётся тот, что опознан по содержимому: расширение
                # приходит от отправителя, и картинка под именем `.pdf`
                # роняла pdf2image, а читаемость объявлялась непроверенной.
                context.uri, pages, file_type=context.file_type,
                # Байты уже прочитаны приёмом: без них каждый растровый
                # документ выкачивался из хранилища второй раз.
                data=self._context_data(context),
            )
        except OcrUnavailable as exc:
            logger.info("Быстрый OCR недоступен для %s: %s", context.uri, exc)
            return None
        except Exception as exc:  # noqa: BLE001 — OCR не обязан быть поднят
            logger.info("Быстрый OCR не отработал на %s: %s", context.uri, exc)
            return None

        return _probe_from_blocks(blocks)

    @staticmethod
    def _context_data(context: DocumentContext) -> Optional[bytes]:
        """Байты документа, если они уже прочитаны. Читать заново незачем."""
        try:
            return context.data
        except Exception as exc:  # noqa: BLE001 — пусть решает провайдер
            logger.debug("Байты %s недоступны: %s", context.uri, exc)
            return None

    def _ocr_sample(self, context: DocumentContext) -> str:
        """Текст последнего быстрого OCR этого документа, если он был."""
        if self._last_probe and self._last_probe[0] == context.uri:
            return self._last_probe[1].text
        return ""

    def _describe(
        self, context: DocumentContext, fingerprint: Optional[int] = None
    ) -> Tuple[Dict[str, Any], Dict[str, Any]]:
        """Маршрут и сигналы приёма. Классификация не вправе ронять вердикт."""
        try:
            routing = self.routing(context)
            classification = context.classification.as_dict()
        except Exception as exc:  # noqa: BLE001
            logger.warning("Классификация %s не удалась: %s", context.uri, exc)
            return {}, {}

        signals: Dict[str, Any] = {
            key: value for key, value in classification.items() if key != "signals"
        }
        # Список чертёжных листов на сотне страниц в отчёт приёма не нужен —
        # доля уже посчитана.
        signals.update({
            key: value for key, value in (classification.get("signals") or {}).items()
            if key != "drawing_pages"
        })
        if fingerprint is None:
            fingerprint = _simhash(_text_sample(context) or self._ocr_sample(context))
        if fingerprint:
            # Нулевой отпечаток — это «текста не было», а не документ. Он
            # уезжал в базу с каждым сканом и вытеснял из выборки почти-дублей
            # настоящие отпечатки.
            signals["fingerprint"] = fingerprint
        return routing, signals

    # ------------------------------- QG-5 классификация и маршрутизация
    def routing(self, context: DocumentContext) -> Dict[str, Any]:
        """
        Тип документа и предложение размещения. Метки доступа проставляются
        здесь, на приёме, но утверждает их человек — предложение уходит в
        очередь модерации, а не применяется само.
        """
        classification = context.classification
        document_type = self._document_type(context, classification)

        # Чертёжные листы поднимают гриф, даже когда их меньшинство: доля
        # ниже 0.5 давала жанр «смешанный», контейнер «Общая документация» и
        # метку «все сотрудники» — приложение с двумя чертежами в договоре на
        # десять листов уезжало в общий доступ.
        try:
            drawing_share = float(classification.signals.get("drawing_share") or 0.0)
        except (TypeError, ValueError):  # pragma: no cover — сигнал чужого вида
            drawing_share = 0.0
        has_drawings = document_type == ROUTE_DRAWING or drawing_share > 0

        return {
            "document_type": document_type,
            "genre": classification.genre,
            "source": classification.source,
            # Растр — это способ хранения, а не род документа; тип теперь
            # задаёт жанр, а признак скана едет отдельным полем.
            "scanned": classification.source == SOURCE_RASTER,
            "drawing_share": round(drawing_share, 3),
            "suggested_container": self._container(document_type),
            "suggested_sensitivity": "confidential" if has_drawings else "internal",
            "suggested_access_labels": (
                ["конструкторский отдел"] if has_drawings
                else self._labels(document_type)
            ),
            "needs_moderation": True,
            "note": "Предложение сформировано системой, утверждает человек.",
        }

    @staticmethod
    def _document_type(context: DocumentContext, classification) -> str:
        """
        Тип задаёт род документа, а не способ его хранения.

        Растровый источник проверялся раньше жанра, и «мультимодальный»
        собирал в себя всё подряд: скан страницы текста, скан ведомости и
        фотография объекта уезжали в один контейнер, хотя жанр листа уже был
        вычислен — то есть считался зря. Теперь «мультимодальный» значит
        ровно то, чем он задуман: растр, род которого определить не удалось.
        """
        if context.kind == filetypes.KIND_CAD or classification.genre == GENRE_DRAWING:
            return ROUTE_DRAWING
        if classification.genre == GENRE_SPREADSHEET:
            return ROUTE_TABULAR
        if classification.genre == GENRE_TEXT:
            return ROUTE_TEXT
        if classification.genre == GENRE_MIXED:
            return ROUTE_MIXED
        if classification.source == SOURCE_RASTER:
            return ROUTE_MULTIMODAL
        return ROUTE_MIXED

    @staticmethod
    def _container(document_type: str) -> str:
        return {
            ROUTE_DRAWING: "Конструкторская документация",
            ROUTE_TABULAR: "Реестры и ведомости",
            ROUTE_TEXT: "Организационные документы",
            ROUTE_MIXED: "Общая документация",
            ROUTE_MULTIMODAL: "Сканы и фотокопии",
        }.get(document_type, "Общая документация")

    @staticmethod
    def _labels(document_type: str) -> List[str]:
        if document_type == ROUTE_DRAWING:
            return ["конструкторский отдел"]
        if document_type == ROUTE_TABULAR:
            return ["планово-экономический отдел"]
        return ["все сотрудники"]


# ===========================================================================
# Помощники
# ===========================================================================

@dataclass
class _OcrProbe:
    """
    Итог быстрого OCR: средняя уверенность, сколько блоков нашлось, какая
    доля текста распознана ниже порога и сам текст пробы.
    """

    confidence: float
    blocks: int
    unreadable_share: float = 0.0
    text: str = ""


def _probe_from_blocks(blocks) -> _OcrProbe:
    """
    Проба по блокам распознавания. Вес блока — длина его текста: строка из
    двух символов не должна весить столько же, сколько абзац.
    """
    if not blocks:
        return _OcrProbe(confidence=0.0, blocks=0)

    weights = [max(1, len(block.text or "")) for block in blocks]
    total = sum(weights)
    confidence = sum(
        block.confidence * weight for block, weight in zip(blocks, weights)
    ) / total
    weak = sum(
        weight for block, weight in zip(blocks, weights)
        if block.confidence < settings.QG_OCR_BLOCK_CONFIDENCE
    )
    sample = " ".join(
        (block.text or "").strip() for block in blocks if (block.text or "").strip()
    )
    return _OcrProbe(
        confidence=confidence,
        blocks=len(blocks),
        unreadable_share=round(weak / total, 3),
        text=sample[:_OCR_SAMPLE_LIMIT],
    )


def _native_text(context: DocumentContext) -> Optional[str]:
    """
    Немного текста из формата, который читается без распознавания.
    `None` — проба неприменима или не удалась, и это не повод для выводов.
    """
    kind = context.kind
    try:
        if kind == filetypes.KIND_PLAIN_TEXT:
            return _decoded(context.data)
        if kind == filetypes.KIND_SPREADSHEET:
            return _spreadsheet_text(context)
        if kind == filetypes.KIND_OFFICE_TEXT:
            return _docx_text(context.data)
    except Exception as exc:  # noqa: BLE001 — проба не обязана удаться
        logger.info("Проба текста %s не удалась: %s", context.uri, exc)
    return None


def _decoded(data: bytes) -> Optional[str]:
    for encoding in ("utf-8", "cp1251"):
        try:
            return data[:_NATIVE_PROBE_LIMIT * 2].decode(encoding)[:_NATIVE_PROBE_LIMIT]
        except UnicodeDecodeError:
            continue
    return None


def _spreadsheet_text(context: DocumentContext) -> Optional[str]:
    """Первые строки книги. CSV читается как текст, xlsx — через openpyxl."""
    if context.file_type == "csv":
        return _decoded(context.data)
    if context.file_type not in ("xlsx", "xlsm"):
        return None

    import io

    import openpyxl
    workbook = openpyxl.load_workbook(io.BytesIO(context.data), read_only=True)
    try:
        parts: List[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(max_row=_NATIVE_PROBE_ROWS, values_only=True):
                parts.extend(str(cell) for cell in row if cell is not None)
            if sum(len(p) for p in parts) >= _MIN_NATIVE_CHARS:
                break
        return " ".join(parts)[:_NATIVE_PROBE_LIMIT]
    finally:
        workbook.close()


def _docx_text(data: bytes) -> Optional[str]:
    """Абзацы и ячейки таблиц документа — ровно до порога пробы."""
    import io

    import docx
    document = docx.Document(io.BytesIO(data))
    parts: List[str] = [p.text for p in document.paragraphs if p.text.strip()]
    if sum(len(p) for p in parts) < _MIN_NATIVE_CHARS:
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells if cell.text.strip())
            if sum(len(p) for p in parts) >= _MIN_NATIVE_CHARS:
                break
    return " ".join(parts)[:_NATIVE_PROBE_LIMIT]


def _can_open(data: bytes, extension: str) -> Tuple[bool, str]:
    """Открывается ли файл своим форматом. Повреждённый отсекается сразу."""
    kind = filetypes.kind_of(extension)
    try:
        if kind == filetypes.KIND_PDF:
            import pymupdf
            document = pymupdf.open(stream=data, filetype="pdf")
            pages = document.page_count
            document.close()
            return (pages > 0, "в PDF нет страниц")
        if kind == filetypes.KIND_IMAGE:
            import io

            from PIL import Image
            image = Image.open(io.BytesIO(data))
            image.verify()
            return True, ""
        if extension == "xlsx":
            import io

            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
            workbook.close()
            return True, ""
        if extension == "docx":
            import io

            import docx
            docx.Document(io.BytesIO(data))
            return True, ""
    except ImportError:  # pragma: no cover — библиотеки может не быть
        return True, ""
    except Exception as exc:  # noqa: BLE001
        return False, str(exc)[:200]
    return True, ""


def _text_sample(context: DocumentContext, limit: int = 4000) -> str:
    """Небольшой кусок текста документа для отпечатка сходства."""
    layer = None
    try:
        layer = context.text_layer
    except Exception:  # noqa: BLE001
        layer = None

    if layer is not None:
        parts: List[str] = []
        for page in sorted(layer.pages)[:3]:
            parts.extend(line.text for line in layer.lines(page))
            if sum(len(p) for p in parts) > limit:
                break
        return " ".join(parts)[:limit]

    if context.kind in (filetypes.KIND_PLAIN_TEXT, filetypes.KIND_SPREADSHEET):
        for encoding in ("utf-8", "cp1251"):
            try:
                return context.data[:limit * 2].decode(encoding)[:limit]
            except UnicodeDecodeError:
                continue
    return ""


def _simhash(text: str) -> int:
    """
    64-битный отпечаток текста. Почти одинаковые документы дают отпечатки,
    различающиеся единицами бит, и это ловит переименованные копии и
    обновлённые редакции, которых точная контрольная сумма не видит.
    """
    tokens = re.findall(r"\w{3,}", (text or "").lower())
    if not tokens:
        return 0

    vector = [0] * _SIMHASH_BITS
    for token in tokens:
        digest = int(hashlib.md5(token.encode("utf-8")).hexdigest()[:16], 16)
        for bit in range(_SIMHASH_BITS):
            vector[bit] += 1 if (digest >> bit) & 1 else -1

    fingerprint = 0
    for bit, value in enumerate(vector):
        if value > 0:
            fingerprint |= 1 << bit
    return fingerprint


def _hamming(left: int, right: int) -> int:
    if not left or not right:
        return _SIMHASH_BITS
    return bin(left ^ right).count("1")


def _document_date(path: str, data: bytes) -> Optional[datetime]:
    """Дата документа: из метаданных формата, иначе из имени файла."""
    extension = filetypes.resolve(path, data).file_type

    if filetypes.kind_of(extension) == filetypes.KIND_PDF:
        try:
            import pymupdf
            document = pymupdf.open(stream=data, filetype="pdf")
            raw = (document.metadata or {}).get("creationDate") or ""
            document.close()
            parsed = _parse_pdf_date(raw)
            if parsed:
                return parsed
        except Exception:  # noqa: BLE001
            pass

    if extension == "docx":
        try:
            import io

            import docx
            created = docx.Document(io.BytesIO(data)).core_properties.created
            if created:
                return created.replace(tzinfo=created.tzinfo or timezone.utc)
        except Exception:  # noqa: BLE001
            pass

    match = _DATE_IN_NAME.search(os.path.basename(path))
    if match:
        try:
            year = int(match.group(0)[:4])
            if 1950 <= year <= datetime.now(timezone.utc).year:
                return datetime(year, 1, 1, tzinfo=timezone.utc)
        except ValueError:
            pass
    return None


def _parse_pdf_date(raw: str) -> Optional[datetime]:
    """PDF хранит дату как `D:20240115120000+03'00'`."""
    match = re.match(r"D:(\d{4})(\d{2})(\d{2})", raw or "")
    if not match:
        return None
    try:
        return datetime(
            int(match.group(1)), int(match.group(2)), int(match.group(3)),
            tzinfo=timezone.utc,
        )
    except ValueError:
        return None
